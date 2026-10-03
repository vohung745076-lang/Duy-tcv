import io
from datetime import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelExportService:
    @staticmethod
    def generate_monthly_candidates_excel(
        candidates_data: List[Dict[str, Any]],
        month: Optional[int] = None,
        year: Optional[int] = None,
    ) -> bytes:
        """
        Tạo luồng byte Excel (.xlsx) danh sách ứng viên theo tháng với đầy đủ:
        - Tên ứng viên
        - Vị trí tuyển dụng
        - Điểm AI đánh giá
        - Trạng thái tuyển dụng (Đã tuyển / Chờ duyệt / Đã loại)
        - Đã gửi mail hoặc Loại (Đã gửi thư mời / Đã gửi thư từ chối / Chưa gửi / Loại)
        - Lý do loại / Ghi chú thẩm định
        - Thời gian & Địa điểm phỏng vấn
        - Người duyệt
        - Email & SĐT
        - Ngày nộp hồ sơ
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DS Ung Vien"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        font_family = "Segoe UI"
        font_title = Font(name=font_family, size=15, bold=True, color="FFFFFF")
        fill_title = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

        font_subtitle = Font(name=font_family, size=10, italic=True, color="475569")

        font_header = Font(name=font_family, size=11, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

        font_data = Font(name=font_family, size=10, color="000000")
        font_bold = Font(name=font_family, size=10, bold=True, color="000000")

        fill_approved = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")  # light green
        fill_rejected = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # light red
        fill_pending = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")   # light yellow
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        font_approved = Font(name=font_family, size=10, bold=True, color="166534")
        font_rejected = Font(name=font_family, size=10, bold=True, color="991B1B")
        font_pending = Font(name=font_family, size=10, bold=True, color="92400E")

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_right = Alignment(horizontal="right", vertical="center")

        thin_border_side = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=thin_border_side,
        )

        headers = [
            "STT",
            "Họ và tên ứng viên",
            "Vị trí ứng tuyển",
            "Điểm AI",
            "Tuyển dụng",
            "Trạng thái Email",
            "Hình thức & Lịch phỏng vấn",
            "Lý do loại / Thẩm định",
            "Người thẩm định",
            "Email liên hệ",
            "Số điện thoại",
            "Ngày nộp hồ sơ",
        ]

        total_cols = len(headers)

        # 1. Title Banner (Row 1)
        period_str = f"THÁNG {month}/{year}" if month and year else (f"NĂM {year}" if year else "TẤT CẢ CÁC THÁNG")
        title_text = f"BÁO CÁO TỔNG HỢP NHÂN SỰ & DANH SÁCH ỨNG VIÊN - {period_str}"
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        title_cell = ws.cell(row=1, column=1, value=title_text)
        title_cell.font = font_title
        title_cell.fill = fill_title
        title_cell.alignment = align_center
        ws.row_dimensions[1].height = 40

        # Fill background for the rest of merged cells in row 1
        for col in range(2, total_cols + 1):
            ws.cell(row=1, column=col).fill = fill_title

        # 2. Subtitle / Metadata (Row 2)
        now_str = datetime.now().strftime("%d/%m/%Y %H:%M")
        subtitle_text = f"Thời điểm xuất file: {now_str} | Tổng số ứng viên: {len(candidates_data)}"
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
        sub_cell = ws.cell(row=2, column=1, value=subtitle_text)
        sub_cell.font = font_subtitle
        sub_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 22

        # Row 3 is blank
        ws.row_dimensions[3].height = 10

        # 3. Header Row (Row 4)
        header_row_idx = 4
        ws.row_dimensions[header_row_idx].height = 28
        for col_idx, header_title in enumerate(headers, 1):
            cell = ws.cell(row=header_row_idx, column=col_idx, value=header_title)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = cell_border

        # 4. Data Rows (Row 5+)
        current_row = 5
        for idx, item in enumerate(candidates_data, 1):
            ws.row_dimensions[current_row].height = 26
            is_even = idx % 2 == 0

            # Raw values
            name = item.get("masked_name") or item.get("original_filename") or f"Ứng viên #{idx}"
            job_title = item.get("job_title") or "Hồ sơ lưu trữ chung"
            score = item.get("overall_score")
            score_text = f"{score}%" if score is not None else "Chưa chấm"

            approval_status = (item.get("approval_status") or "PENDING").upper()
            if approval_status == "APPROVED":
                tuyen_dung_text = "Đã tuyển dụng / Phù hợp"
                tuyen_dung_font = font_approved
                tuyen_dung_fill = fill_approved
            elif approval_status == "REJECTED":
                tuyen_dung_text = "Đã loại / Từ chối"
                tuyen_dung_font = font_rejected
                tuyen_dung_fill = fill_rejected
            else:
                tuyen_dung_text = "Chờ HR duyệt"
                tuyen_dung_font = font_pending
                tuyen_dung_fill = fill_pending

            # Email sent status / Loại
            interview_type = item.get("interview_type")
            interview_time = item.get("interview_time")
            interview_location = item.get("interview_location")
            rejection_reason = item.get("rejection_reason") or ""

            if approval_status == "APPROVED" and interview_time:
                email_status_text = f"Đã gửi thư mời ({interview_type or 'Phỏng vấn'})"
                email_status_font = font_approved
                email_status_fill = fill_approved
            elif approval_status == "REJECTED":
                email_status_text = "Đã gửi thư từ chối / Loại"
                email_status_font = font_rejected
                email_status_fill = fill_rejected
            elif approval_status == "APPROVED":
                email_status_text = "Đã duyệt (Chưa gửi mail)"
                email_status_font = font_pending
                email_status_fill = fill_pending
            else:
                email_status_text = "Chưa gửi mail"
                email_status_font = font_data
                email_status_fill = fill_zebra if is_even else None

            # Schedule text
            if interview_time or interview_location:
                schedule_text = f"{interview_type or 'PV'} | {interview_time or 'Chưa định giờ'} | {interview_location or 'Online'}"
            else:
                schedule_text = "-"

            reason_text = rejection_reason if rejection_reason else ("Phù hợp tiêu chí JD" if approval_status == "APPROVED" else "-")
            reviewer = item.get("reviewed_by") or "Hệ thống AI / Chờ HR"
            email = item.get("email") or "-"
            phone = item.get("phone") or "-"

            created_at_raw = item.get("created_at") or ""
            if created_at_raw:
                try:
                    dt = datetime.fromisoformat(created_at_raw.replace("Z", "+00:00"))
                    created_at_text = dt.strftime("%d/%m/%Y %H:%M")
                except Exception:
                    created_at_text = str(created_at_raw)[:16]
            else:
                created_at_text = "-"

            row_values = [
                (idx, align_center, font_data, fill_zebra if is_even else None),
                (name, align_left, font_bold, fill_zebra if is_even else None),
                (job_title, align_left, font_data, fill_zebra if is_even else None),
                (score_text, align_center, font_bold, fill_zebra if is_even else None),
                (tuyen_dung_text, align_center, tuyen_dung_font, tuyen_dung_fill),
                (email_status_text, align_center, email_status_font, email_status_fill),
                (schedule_text, align_left, font_data, fill_zebra if is_even else None),
                (reason_text, align_left, font_data, fill_zebra if is_even else None),
                (reviewer, align_left, font_data, fill_zebra if is_even else None),
                (email, align_left, font_data, fill_zebra if is_even else None),
                (phone, align_center, font_data, fill_zebra if is_even else None),
                (created_at_text, align_center, font_data, fill_zebra if is_even else None),
            ]

            for col_idx, (val, alignment, font_style, fill_style) in enumerate(row_values, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.alignment = alignment
                cell.font = font_style
                if fill_style:
                    cell.fill = fill_style
                cell.border = cell_border

            current_row += 1

        # 5. Auto-adjust Column Widths
        min_widths = {
            1: 7,   # STT
            2: 24,  # Họ và tên
            3: 26,  # Vị trí
            4: 12,  # Điểm AI
            5: 22,  # Tuyển dụng
            6: 24,  # Trạng thái Email
            7: 28,  # Lịch PV
            8: 26,  # Lý do loại
            9: 22,  # Người thẩm định
            10: 25, # Email
            11: 15, # Phone
            12: 18, # Ngày nộp
        }

        for col_idx in range(1, total_cols + 1):
            col_letter = get_column_letter(col_idx)
            max_len = min_widths.get(col_idx, 15)
            # Find max length of contents in this column (from row 4 downwards)
            for row in range(4, current_row):
                cell_val = ws.cell(row=row, column=col_idx).value
                if cell_val is not None:
                    # rough length calculation for Vietnamese text
                    val_str = str(cell_val)
                    max_len = max(max_len, min(len(val_str) + 4, 45))
            ws.column_dimensions[col_letter].width = max_len

        # Save to buffer
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()


excel_export_service = ExcelExportService()
